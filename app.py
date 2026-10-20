import hashlib
import io

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font

st.set_page_config(page_title="Portalista Report", page_icon="📊", layout="wide")

st.markdown(
    """
    <style>
        :root { --portal-navy: #0f2742; --portal-slate: #536477; --portal-teal: #0f766e; }
        .stApp { background: #f6f8fb; color: var(--portal-navy); }
        [data-testid="stHeader"] { background: rgba(246, 248, 251, 0.92); }
        [data-testid="stMainBlockContainer"] { max-width: 1180px; padding-top: 2rem; padding-bottom: 3.5rem; }
        .portal-hero { background: linear-gradient(120deg, #102b49, #164d60); border-radius: 18px; color: #fff; padding: 2.1rem 2.25rem; margin-bottom: 1.5rem; }
        .portal-kicker { color: #8ee2dd; font-size: 0.8rem; font-weight: 700; letter-spacing: 0.14em; margin-bottom: 0.4rem; }
        .portal-hero h1 { color: #fff; font-size: 2rem; margin: 0; }
        .portal-hero p { color: #d6e4eb; margin: 0.55rem 0 0; }
        .step-card { background: #fff; border: 1px solid #dbe4eb; border-radius: 14px; min-height: 112px; padding: 1rem 1.1rem; }
        .step-number { color: var(--portal-teal); font-size: 0.8rem; font-weight: 700; letter-spacing: 0.08em; }
        .step-card h3 { color: var(--portal-navy); font-size: 1rem; margin: 0.35rem 0; }
        .step-card p { color: var(--portal-slate); font-size: 0.9rem; margin: 0; }
        [data-testid="stFileUploader"] { background: #fff; border: 1px solid #dbe4eb; border-radius: 14px; padding: 0.75rem 0.85rem; }
        [data-testid="stMetric"] { background: #fff; border: 1px solid #dbe4eb; border-radius: 14px; padding: 0.9rem 1rem; }
        [data-testid="stBaseButton-primary"]:not(:disabled), [data-testid="stDownloadButton"] > button { background: var(--portal-teal); border-color: var(--portal-teal); color: white; }
        [data-testid="stBaseButton-primary"]:not(:disabled):hover { background: #115e59; border-color: #115e59; }
        [data-testid="stDownloadButton"] > button { width: 100%; }
    </style>
    """,
    unsafe_allow_html=True,
)

THAI_MONTHS = ["", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
               "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]

REQUIRED_COLUMNS = ["building", "policy", "session time", "download", "upload", "start"]


def find_col(df, keyword):
    """ค้นหาคอลัมน์ที่มีคำ keyword ในชื่อ (ไม่สนตัวพิมพ์เล็ก-ใหญ่)"""
    for col in df.columns:
        if keyword in str(col).lower():
            return col
    return None


def to_hours(text):
    if pd.isna(text):
        return 0.0
    h, m, s = str(text).split(":")
    return int(h) + int(m) / 60 + int(s) / 3600


def display_name(building):
    """แปลงชื่อสถานที่ดิบเป็นชื่อเต็มสำหรับแสดงผล"""
    if building.startswith("สพ."):
        return "สถานพินิจฯ " + building[len("สพ."):]
    if building == "กรมพินิจและคุ้มครองเด็กและเยาวชน":
        return "สถานพินิจฯ กรมพินิจและคุ้มครองเด็กฯ"
    return building


def build_summary(session_df):
    building_col = find_col(session_df, "building")
    policy_col = find_col(session_df, "policy")
    time_col = find_col(session_df, "session time")
    download_col = find_col(session_df, "download")
    upload_col = find_col(session_df, "upload")
    start_col = find_col(session_df, "start")

    df = session_df.copy()
    df["_hours"] = df[time_col].apply(to_hours)
    df["_volume"] = df[download_col].fillna(0) + df[upload_col].fillna(0)

    rows = []
    for building, group in df.groupby(building_col, sort=True):
        rows.append({
            "สถานที่": display_name(str(building)),
            "ประชาชนผู้มารับบริการ (ครั้ง)": int((group[policy_col] == "Guest").sum()),
            "เจ้าหน้าที่กรมพินิจฯ (ครั้ง)": int((group[policy_col] == "Staff").sum()),
            "ชั่วโมงการใช้งาน": group["_hours"].sum(),
            "ปริมาณการใช้งาน": group["_volume"].sum()
        })
    summary = pd.DataFrame(rows)

    total_row = pd.DataFrame([{
        "สถานที่": "รวมทั้งหมด",
        "ประชาชนผู้มารับบริการ (ครั้ง)": summary["ประชาชนผู้มารับบริการ (ครั้ง)"].sum(),
        "เจ้าหน้าที่กรมพินิจฯ (ครั้ง)": summary["เจ้าหน้าที่กรมพินิจฯ (ครั้ง)"].sum(),
        "ชั่วโมงการใช้งาน": summary["ชั่วโมงการใช้งาน"].sum(),
        "ปริมาณการใช้งาน": summary["ปริมาณการใช้งาน"].sum()
    }])
    summary = pd.concat([summary, total_row], ignore_index=True)

    month_num = pd.to_datetime(df[start_col]).dt.month.mode()[0]
    title = f"จำนวน ผู้ใช้งาน แยกตาม หน่วยงานกรมพินิจฯ ประจำเดือน{THAI_MONTHS[month_num]}"

    return summary, title


def build_excel(internet_df, session_df, summary, title):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        internet_df.to_excel(writer, index=False, sheet_name="internet user")
        session_df.to_excel(writer, index=False, sheet_name="session logs")

        ws = writer.book.create_sheet("สรุป")

        ws.append([title, None, None, None, None])
        ws.merge_cells("A1:E1")
        ws["A1"].font = Font(bold=True)
        ws["A1"].alignment = Alignment(horizontal="center")

        headers = list(summary.columns)
        ws.append(headers)
        for cell in ws[2]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        last_row_idx = len(summary) - 1
        for i, row in summary.iterrows():
            ws.append(list(row))
            r = ws.max_row
            is_total = (i == last_row_idx)
            for col_idx in range(1, 6):
                cell = ws.cell(row=r, column=col_idx)
                if is_total:
                    cell.font = Font(bold=True)
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="left")
                elif col_idx in (2, 3):
                    cell.alignment = Alignment(horizontal="center")
                    cell.number_format = "#,##0"
                elif col_idx == 4:
                    cell.alignment = Alignment(horizontal="center")
                    cell.number_format = "#,##0.00"
                elif col_idx == 5:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0.00"

        ws.column_dimensions["A"].width = 50
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 20
        ws.column_dimensions["E"].width = 22

    buffer.seek(0)
    return buffer


def file_signature(uploaded_file):
    if uploaded_file is None:
        return None
    return (uploaded_file.name, uploaded_file.size, hashlib.sha256(uploaded_file.getvalue()).hexdigest())


def clear_report_if_inputs_changed(current_signature):
    previous_signature = st.session_state.get("input_signature")
    if previous_signature is not None and previous_signature != current_signature:
        for key in ("report_summary", "report_title", "report_excel", "input_signature"):
            st.session_state.pop(key, None)


st.markdown(
    """
    <section class="portal-hero">
      <div class="portal-kicker">PORTALISTA / REPORT</div>
      <h1>รายงานการใช้งานอินเทอร์เน็ต</h1>
      <p>สรุปการใช้งานแยกตามสถานที่จากไฟล์ Internet Users และ Session Logs</p>
    </section>
    """,
    unsafe_allow_html=True,
)

step_columns = st.columns(3)
steps = [
    ("01 / UPLOAD", "อัปโหลดไฟล์", "เลือกไฟล์ Excel ทั้ง 2 รายการ"),
    ("02 / GENERATE", "สร้างรายงาน", "ตรวจสอบข้อมูลและสรุปผลตามสถานที่"),
    ("03 / DOWNLOAD", "ดาวน์โหลด Excel", "รับรายงานพร้อมข้อมูลต้นทางและสรุป"),
]
for column, (number, heading, detail) in zip(step_columns, steps):
    with column:
        st.markdown(
            f'<section class="step-card"><div class="step-number">{number}</div><h3>{heading}</h3><p>{detail}</p></section>',
            unsafe_allow_html=True,
        )

st.markdown("### 1. อัปโหลดข้อมูล")
st.caption("รองรับเฉพาะไฟล์ .xlsx เพื่อให้ระบบอ่านข้อมูลด้วยรูปแบบเดียวกัน")
upload_columns = st.columns(2, gap="large")
with upload_columns[0]:
    st.markdown("#### Internet Users")
    st.caption("ไฟล์ข้อมูลผู้ใช้งานอินเทอร์เน็ตต้นทาง")
    file1 = st.file_uploader("เลือกไฟล์ Internet Users", type=["xlsx"], key="internet_users_file")
    if file1:
        st.success("ไฟล์ Internet Users พร้อมใช้งาน")
    else:
        st.info("รอไฟล์ Internet Users")
with upload_columns[1]:
    st.markdown("#### Session Logs")
    st.caption("ไฟล์บันทึก session สำหรับคำนวณรายงาน")
    file2 = st.file_uploader("เลือกไฟล์ Session Logs", type=["xlsx"], key="session_logs_file")
    if file2:
        st.success("ไฟล์ Session Logs พร้อมใช้งาน")
    else:
        st.info("รอไฟล์ Session Logs")

current_signature = (file_signature(file1), file_signature(file2))
clear_report_if_inputs_changed(current_signature)
files_ready = file1 is not None and file2 is not None

with st.expander("ตรวจสอบโครงสร้างข้อมูลก่อนสร้างรายงาน"):
    st.write("ไฟล์ Session Logs ต้องมีคอลัมน์ที่ชื่อมีคำต่อไปนี้: " + ", ".join(REQUIRED_COLUMNS))
    st.write("ไฟล์ที่ดาวน์โหลดจะมี 3 แผ่นงาน: `internet user`, `session logs` และ `สรุป`.")
    st.caption("ข้อมูลจะถูกประมวลผลโดยเซิร์ฟเวอร์ที่รัน Streamlit เพื่อสร้างรายงาน")

st.markdown("### 2. สร้างรายงาน")
if not files_ready:
    st.caption("อัปโหลดไฟล์ให้ครบ 2 รายการก่อนจึงจะสร้างรายงานได้")

if st.button("สร้างรายงาน", type="primary", disabled=not files_ready, use_container_width=True):
    try:
        with st.spinner("กำลังตรวจสอบข้อมูลและสร้างรายงาน..."):
            internet_df = pd.read_excel(file1, engine="openpyxl")
            session_df = pd.read_excel(file2, engine="openpyxl")

            if internet_df.empty:
                raise ValueError("ไฟล์ Internet Users ไม่มีแถวข้อมูล")
            if session_df.empty:
                raise ValueError("ไฟล์ Session Logs ไม่มีแถวข้อมูล")

            missing = [column for column in REQUIRED_COLUMNS if find_col(session_df, column) is None]
            if missing:
                raise ValueError("ไม่พบคอลัมน์ที่ต้องใช้ในไฟล์ Session Logs: " + ", ".join(missing))

            building_col = find_col(session_df, "building")
            start_col = find_col(session_df, "start")
            if session_df[building_col].dropna().empty or pd.to_datetime(session_df[start_col], errors="coerce").dropna().empty:
                raise ValueError("ไฟล์ Session Logs ไม่มีแถวข้อมูลที่ใช้สร้างรายงานได้")

            summary, title = build_summary(session_df)
            excel_data = build_excel(internet_df, session_df, summary, title).getvalue()

        st.session_state.report_summary = summary
        st.session_state.report_title = title
        st.session_state.report_excel = excel_data
        st.session_state.input_signature = current_signature
        st.success("สร้างรายงานเรียบร้อยแล้ว")
    except ValueError as error:
        st.error(str(error))
    except Exception:
        st.error("ไม่สามารถอ่านหรือประมวลผลไฟล์ได้ กรุณาตรวจสอบว่าเป็นไฟล์ .xlsx ที่ไม่เสียหาย และรูปแบบข้อมูลถูกต้อง")

if "report_summary" in st.session_state:
    summary = st.session_state.report_summary
    total = summary.iloc[-1]
    st.markdown("### 3. ผลลัพธ์และดาวน์โหลด")
    st.success(st.session_state.report_title)
    metric_columns = st.columns(4)
    metric_columns[0].metric("สถานที่", f"{max(len(summary) - 1, 0):,}")
    metric_columns[1].metric("ประชาชนผู้มารับบริการ (ครั้ง)", f"{int(total['ประชาชนผู้มารับบริการ (ครั้ง)']):,}")
    metric_columns[2].metric("เจ้าหน้าที่กรมพินิจฯ (ครั้ง)", f"{int(total['เจ้าหน้าที่กรมพินิจฯ (ครั้ง)']):,}")
    metric_columns[3].metric("ชั่วโมงการใช้งาน", f"{total['ชั่วโมงการใช้งาน']:,.2f}")

    st.dataframe(
        summary,
        use_container_width=True,
        hide_index=True,
        column_config={
            "ประชาชนผู้มารับบริการ (ครั้ง)": st.column_config.NumberColumn(format="%,d"),
            "เจ้าหน้าที่กรมพินิจฯ (ครั้ง)": st.column_config.NumberColumn(format="%,d"),
            "ชั่วโมงการใช้งาน": st.column_config.NumberColumn(format="%,.2f"),
            "ปริมาณการใช้งาน": st.column_config.NumberColumn(format="%,.2f"),
        },
    )
    st.download_button(
        label="ดาวน์โหลดรายงาน Excel",
        data=st.session_state.report_excel,
        file_name="location_report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )
else:
    st.info("ผลลัพธ์จะแสดงที่นี่หลังจากสร้างรายงานสำเร็จ")
